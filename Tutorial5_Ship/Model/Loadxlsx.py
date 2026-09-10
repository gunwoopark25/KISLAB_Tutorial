import openpyxl

class Loadxlsx:
    # InputData 주소
    xlsx_PATH = r"C:\Users\User\SynologyDrive\01. Code 공부\KISLAB_Tutorial\Tutorial5_Ship\Input\320K offset.xlsx"

    @staticmethod
    def load(path=None):
        # 절대 주소에서 파일 불러오기
        path = Loadxlsx.xlsx_PATH

        # data_only = True : 데이터에서의 수식 결과값을 불러오겠다는 의미
        workbook = openpyxl.load_workbook(path, data_only=True)
        sheet = workbook.worksheets[0]

        WaterLine_row = next(sheet.iter_rows(min_row=3, max_row=3, values_only=True))

        # Bottom Line을 0.0으로 미리 정의
        waterlines = [0.0]  

        # WaterLine_row에서 숫자Data만 가져오는 과정
        for value in WaterLine_row[2:]:
            if not isinstance(value, (int, float)):
                break
            waterlines.append(float(value))

        # 반폭 OffsetData를 Dictionary로 저장
        half_breadth = {}

        # 7부터하는 이유 : 5~6까지는 -0.333, -0.166이고 Station No가 0 계산을 시작하기 위해서
        for row in sheet.iter_rows(min_row=7, values_only=True):
            station_value = row[0]

            # 숫자가 아닌 행(빈 행 등)은 건너뛴다
            if not isinstance(station_value, (int, float)):
                continue

            station = int(station_value)
            # mm단위를 m로 변경
            offsets_mm = row[1:1 + len(waterlines)]
            offsets_m = []

            for value_mm in offsets_mm:
                if not isinstance(value_mm, (int, float)):
                    raise ValueError(f"Station {station}의 반폭 값이 숫자가 아닙니다: {value_mm!r}")
                offsets_m.append(value_mm / 1000)

            half_breadth[station] = offsets_m

        stations = sorted(half_breadth.keys())

        return {
            "stations": stations,
            "waterlines": waterlines,
            "half_breadth": half_breadth,
        }

    # (xs, ys) 데이터에서 x에 해당하는 y를 선형보간으로 구한다
    @staticmethod
    def linear_interpolation(xs, ys, x):
        for i in range(len(xs) - 1):
            if xs[i] <= x <= xs[i + 1]:
                ratio = (x - xs[i]) / (xs[i + 1] - xs[i])
                return ys[i] + (ys[i + 1] - ys[i]) * ratio

        raise ValueError(f"{x}는 보간 범위({xs[0]} ~ {xs[-1]})를 벗어났습니다.")

    # 원본 waterline은 0~16m가 1m 간격, 16~30m가 2m 간격이라
    # Simpson 법칙을 그대로 쓸 수 없다.
    # 빠져있는 17,19,...,29m를 선형보간해서 0~30m 전체를 1m 등간격으로 만든다.
    @staticmethod
    def interpolate(data):
        raw_waterlines = data["waterlines"]
        raw_half_breadth = data["half_breadth"]

        # 0, 1, 2, ..., 30 (31개)
        waterlines = [float(z) for z in range(int(raw_waterlines[-1]) + 1)]

        half_breadth = {}

        for station, raw_offsets in raw_half_breadth.items():
            offsets = []

            for z in waterlines:
                offsets.append(Loadxlsx.linear_interpolation(raw_waterlines, raw_offsets, z))

            half_breadth[station] = offsets

        return {
            "stations": data["stations"],
            "waterlines": waterlines,
            "half_breadth": half_breadth,
        }
